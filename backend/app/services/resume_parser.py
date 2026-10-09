import os
from pathlib import Path
from typing import Tuple
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from docx import Document

class ResumeParserError(Exception):
    """Base exception for resume parsing errors"""
    pass

class ScannedPdfError(ResumeParserError):
    """Exception raised when a PDF contains only scanned images and no selectable text"""
    pass

class EmptyFileError(ResumeParserError):
    """Exception raised when an uploaded file contains no extractable text"""
    pass

class UnsupportedFormatError(ResumeParserError):
    """Exception raised when the file extension or signature is unsupported"""
    pass

def validate_resume_file(filename: str, file_bytes: bytes, max_size_mb: int = 10) -> str:
    """
    Validates file extension, size, and header signatures.
    Returns detected extension ('pdf', 'docx', or 'txt').
    """
    if len(file_bytes) == 0:
        raise EmptyFileError("Uploaded file is empty (0 bytes).")

    max_bytes = max_size_mb * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise ResumeParserError(f"File size exceeds maximum limit of {max_size_mb} MB.")

    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in ["pdf", "docx", "txt"]:
        raise UnsupportedFormatError(f"Unsupported file extension '.{ext}'. Supported formats are PDF, DOCX, and TXT.")

    # Validate header signature
    if ext == "pdf":
        if not file_bytes.startswith(b"%PDF-"):
            raise UnsupportedFormatError("File does not appear to be a valid PDF format.")
    elif ext == "docx":
        # ZIP magic numbers PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise UnsupportedFormatError("File does not appear to be a valid DOCX package.")
    elif ext == "txt":
        # Check text is decodable
        try:
            file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                file_bytes.decode("latin-1")
            except Exception:
                raise UnsupportedFormatError("File does not appear to be readable plain text.")

    return ext

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using PyMuPDF (fitz).
    Detects scanned PDFs with no text.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ResumeParserError(f"Failed to open PDF document: {str(e)}")

    extracted_pages = []
    total_images = 0

    for page_idx in range(len(doc)):
        page = doc[page_idx]
        text = page.get_text("text").strip()
        if text:
            extracted_pages.append(text)
        total_images += len(page.get_images())

    full_text = "\n\n".join(extracted_pages).strip()

    if not full_text:
        if total_images > 0:
            raise ScannedPdfError(
                "PDF contains scanned images with no selectable text. "
                "Optical Character Recognition (OCR) is required to process this document."
            )
        raise EmptyFileError("PDF document contains no readable text.")

    return full_text

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX bytes using python-docx.
    """
    import io
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ResumeParserError(f"Failed to open DOCX document: {str(e)}")

    paragraphs = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n".join(paragraphs).strip()
    if not full_text:
        raise EmptyFileError("DOCX document contains no readable text.")

    return full_text

def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Extracts text from plain text (TXT) file bytes with encoding fallbacks.
    """
    try:
        text = file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1").strip()
        except Exception as e:
            raise ResumeParserError(f"Failed to decode text document: {str(e)}")

    if not text:
        raise EmptyFileError("Text document contains no readable content.")

    return text

def parse_resume_content(filename: str, file_bytes: bytes, max_size_mb: int = 10) -> Tuple[str, str]:
    """
    Validates and extracts plain text from PDF, DOCX, or TXT file bytes.
    Returns (extracted_text, file_type).
    """
    file_type = validate_resume_file(filename, file_bytes, max_size_mb)
    if file_type == "pdf":
        text = extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        text = extract_text_from_docx(file_bytes)
    elif file_type == "txt":
        text = extract_text_from_txt(file_bytes)
    else:
        raise UnsupportedFormatError(f"Unsupported format: {file_type}")

    return text, file_type
